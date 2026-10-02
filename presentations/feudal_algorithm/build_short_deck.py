#!/usr/bin/env python3
"""Build a matched PowerPoint / PDF deck, with editable text and diagrams.

Requires python-pptx, reportlab, matplotlib and Pillow. Equations are rendered
at 240 dpi; their editable LaTeX is retained in the source and speaker notes.
"""
from __future__ import annotations

import hashlib
import math
import os
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/feudal-slides-matplotlib")
import matplotlib
matplotlib.use("Agg")
from matplotlib.mathtext import math_to_image
from matplotlib.font_manager import FontProperties
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "equations"
ASSETS.mkdir(exist_ok=True)
W, H = 960, 540
BG, INK, MUTED = "FFFFFF", "20262D", "5A6570"
TEAL, PURPLE, AMBER = "395D80", "395D80", "395D80"
PALE, LAVENDER, GOLD, LINE = "FFFFFF", "FFFFFF", "FFFFFF", "DCE1E6"
WHITE = "FFFFFF"
FONT = "DejaVu Sans"
FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
if not FONT_DIR.exists():
    FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
for suffix, filename in [("", "DejaVuSans.ttf"), ("-Bold", "DejaVuSans-Bold.ttf")]:
    pdfmetrics.registerFont(TTFont("Deck" + suffix, str(FONT_DIR / filename)))

SLIDES = []

def wrap(value, width, size, bold=False):
    font = "Deck-Bold" if bold else "Deck"
    lines = []
    for paragraph in value.split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = (line + " " + word).strip()
            if pdfmetrics.stringWidth(candidate, font, size) > width and line:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines

class Slide:
    def __init__(self, section, title, takeaway, notes, source, appendix=False):
        self.section, self.title, self.notes, self.source = section, title, notes, source
        self.elements, self.equations = [], []
        self.number = len(SLIDES) + 1
        self.appendix = appendix
        SLIDES.append(self)
        self.rect(0, 0, W, H, BG)
        self.text(section.upper(), 42, 25, 870, 11, TEAL, True)
        self.text(title, 42, 50, 880, 31, INK, True, max_h=77)
        if takeaway:
            self.line(42, 467, 918, 467, LINE, 1)
            self.text(takeaway, 56, 477, 846, 15, TEAL, True, max_h=24)
        self.text(source, 42, 518, 825, 8.2, MUTED)
        self.text(f"{self.number:02d}", 874, 516, 44, 11, MUTED, align="right")

    def rect(self, x, y, w, h, fill=WHITE, stroke=None, radius=0):
        self.elements.append(dict(kind="rect", x=x,y=y,w=w,h=h,fill=fill,stroke=stroke,radius=radius))

    def text(self, value, x, y, w, size=20, color=INK, bold=False, align="left", max_h=None):
        lines = wrap(value,w,size,bold)
        height = len(lines)*size*1.18 + 5
        if max_h is not None and height > max_h:
            raise ValueError(f"Slide {self.number}: text too tall ({height:.1f} > {max_h}): {value}")
        if y + height > H + 3:
            raise ValueError(f"Slide {self.number}: text off page: {value}")
        self.elements.append(dict(kind="text",text="\n".join(lines),x=x,y=y,w=w,h=height,size=size,color=color,bold=bold,align=align))
        return height

    def eq(self, latex, x, y, w=860, h=58, size=27, color=INK):
        key = hashlib.sha256((latex+str(size)+color).encode()).hexdigest()[:18]
        path = ASSETS / f"{key}.png"
        if not path.exists():
            math_to_image("$"+latex+"$",str(path),prop=FontProperties(size=size),dpi=240,format="png",color="#"+color)
        iw,ih=Image.open(path).size
        natural_w,natural_h=iw*72/240,ih*72/240
        scale=min(w/natural_w,h/natural_h,1.0)
        dw,dh=natural_w*scale,natural_h*scale
        self.elements.append(dict(kind="image",path=path,x=x,y=y+(h-dh)/2,w=dw,h=dh))
        self.equations.append(latex)

    def line(self, x1,y1,x2,y2,color=LINE,width=2,arrow=False):
        self.elements.append(dict(kind="line",x1=x1,y1=y1,x2=x2,y2=y2,color=color,width=width))
        if arrow:
            a=math.atan2(y2-y1,x2-x1)
            points=[(x2,y2),(x2-9*math.cos(a)+4*math.sin(a),y2-9*math.sin(a)-4*math.cos(a)),(x2-9*math.cos(a)-4*math.sin(a),y2-9*math.sin(a)+4*math.cos(a))]
            self.elements.append(dict(kind="polygon",points=points,fill=color))

    def circle(self,x,y,r,fill):
        self.elements.append(dict(kind="circle",x=x-r,y=y-r,w=2*r,h=2*r,fill=fill))

    def card(self, x,y,w,h,label,body,color=TEAL,fill=WHITE,body_size=20):
        self.rect(x,y,w,h,fill,LINE,0)
        self.text(label,x+18,y+16,w-36,12,color,True)
        self.text(body,x+18,y+44,w-36,body_size,max_h=h-49)

    def bullets(self, items, x=48,y=140,w=860,size=21,gap=16,color=TEAL):
        for item in items:
            self.circle(x+4,y+10,3,color)
            height=self.text(item,x+20,y,w-20,size)
            y+=height+gap

    def table(self, headers, rows, widths, x=42,y=140,row_h=48,size=17):
        total=sum(widths)
        self.rect(x,y,total,35,INK,radius=4)
        px=x
        for txt,ww in zip(headers,widths):
            self.text(txt,px+12,y+10,ww-24,12,WHITE,True)
            px+=ww
        for i,row in enumerate(rows):
            yy=y+35+i*row_h
            self.rect(x,yy,total,row_h,WHITE if i%2==0 else "EDF1F6")
            px=x
            for txt,ww in zip(row,widths):
                self.text(txt,px+12,yy+10,ww-24,size,max_h=row_h-8)
                px+=ww


# 01 / Architecture
s=Slide("Feudal MAPPO / the main algorithm", "Two levels, one environment", "The worker is shared across agents; the manager is centralized even at execution.",
"Allow about 8–10 minutes for this eight-slide walkthrough. This is the current repository implementation, not a generic textbook description. At each step the manager reads global state X and constructs a latent s_i and a directional goal g_i for every agent. The default manager core is an MLP. The worker receives its local observation and goal conditioning, but the goals carry centralized information. One policy is shared across agents. Centralized critics read X. Continuous policies produce Gaussian parameters; discrete policies produce masked categorical logits. Subscripts t and i denote timestep and agent; theta denotes manager parameters and phi denotes worker parameters. The goal-conditioning vector u is pooled from recent goals. A step is one decision in the selected environment. No performance ranking or variant survey is included.",
"Sources: manager.py::FeudalManager • worker.py::FeudalWorker • trainer.py::_env_step")
s.card(44,150,192,115,"GLOBAL STATE", "Xₜ: joint information",MUTED,WHITE,19)
s.card(279,150,269,115,"CENTRAL MANAGER", "θ → latents sₜ,ᵢ\nand goals gₜ,ᵢ",PURPLE,WHITE,22)
s.card(594,150,324,115,"SHARED WORKER", "πφ(aₜ,ᵢ | oₜ,ᵢ, uₜ,ᵢ)",TEAL,WHITE,22)
s.line(239,208,272,208,TEAL,2,True)
s.line(551,208,587,208,TEAL,2,True)
s.text("pool goals",554,276,115,13,TEAL)
s.card(44,318,504,104,"CENTRALIZED CRITICS", "Predict extrinsic, manager and optional intrinsic returns.",MUTED,WHITE,19)
s.card(594,318,324,104,"LOCAL OBSERVATION", "oₜ,ᵢ + goal conditioning uₜ,ᵢ",TEAL,WHITE,19)
s.line(756,313,756,273,TEAL,2,True)

# 02 / Goal semantics
s=Slide("Goals / timing and meaning", "A new goal every step; progress over c steps", "The horizon controls pooling and evaluation. The default is c = 10 steps.",
"The base manager emits a normalized goal every step, not once every c steps. Each goal describes a direction in a learned latent space, not a destination or travel distance. The worker takes the normalized sum of up to c recent goals from its own episode. K_t contains available offsets 0,...,c-1 in the same episode and rollout. M_theta,i denotes the manager goal generator, which reads all agents' latents s_t. The default encoder is centralized: (s_t,1,...,s_t,N)=f_theta(X_t). The manager later evaluates the goal from t against the displacement s_(t+c,i)-s_(t,i). U is stabilized unit normalization: U(v)=v/(||v||_2+epsilon_g), epsilon_g=1e-6. Thus zero stays zero and nonzero vectors are approximately unit length. The two-layer tanh worker concatenates local observation with u by default.",
"Sources: manager.py::goal_channel / transition_cosine • worker.py::FeudalWorker")
s.eq(r"g_{t,i}=U(M_{\theta,i}(s_t)),\qquad u_{t,i}=U\!\left(\sum_{k\in K_t}g_{t-k,i}\right)",48,133,870,62)
for j,label in enumerate(["t − c + 1","…","t − 1","t","…","t + c"]):
    xx=95+j*151
    s.circle(xx,274,5,TEAL)
    s.text(label,xx-45,291,90,18,align="center")
    if j<5:s.line(xx+8,274,xx+143,274,LINE,2)
s.line(95,239,548,239,TEAL,2)
s.text("Recent goals condition this action",86,208,510,18,TEAL,True)
s.line(548,348,850,348,TEAL,2,True)
s.text("Future displacement scores gₜ",547,364,366,18,TEAL,True)
s.text("U: unit normalization   •   i: agent   •   s: latent state   •   Kₜ: valid recent offsets",48,424,862,16,MUTED)

# 03 / Objectives
s=Slide("Objectives / with and without intrinsic reward", "Task success and goal following", "The manager always uses extrinsic advantage; α adds goal following to the worker’s objective.",
"Without intrinsic reward, alpha=0 and the worker uses normalized extrinsic advantage only. Goals still condition the policy, but no explicit reward asks the worker to obey them. The manager still learns task-weighted transition alignment. With intrinsic reward, each stream has a separate critic and GAE. Normalize each advantage separately over rollout time, per environment and agent, then combine. This stops raw reward scale from deciding the mixture. A^W is the advantage passed to worker PPO. E denotes worker extrinsic and I denotes worker intrinsic; hats denote standardized advantages. Alpha weights standardized signals, not raw rewards and not an exact fraction of the resulting gradient norm. At configured alpha=0 the intrinsic critic and reward computation are omitted. Extrinsic reward is the environment's learning reward: usually a shared team reward, or per-agent rewards when configured. The intrinsic reward is not added to the environment reward before GAE. The manager has its own extrinsic critic and never uses the intrinsic advantage.",
"Sources: mappo.py::ppo_update / manager_update")
s.card(44,141,405,116,"EXTRINSIC STREAM", "Task reward → critic → GAE\n→ normalized advantage Âᴱ",TEAL,WHITE,21)
s.card(513,141,405,116,"INTRINSIC STREAM", "Goal progress → critic → GAE\n→ normalized advantage Âᴵ",TEAL,WHITE,21)
s.eq(r"A^W_{t,i}=\widehat A^E_{t,i}+\alpha\widehat A^I_{t,i}",196,285,660,76,35)
s.text("α = 0",64,382,350,24,INK,True)
s.text("Learn from task reward only.",64,416,378,19,MUTED)
s.text("α > 0",531,382,350,24,INK,True)
s.text("Also learn to follow the goals.",531,416,378,19,MUTED)

# 04 / Intrinsic reward
s=Slide("Intrinsic reward / what the worker is paid for", "Score the action’s actual successor", "Alignment is measured after the action, using the successor state before any reset.",
"This equation is the transition-aligned intrinsic reward wired into trainer._apply_intrinsic_reward. s_t,i^+ is the latent of the successor produced by action a_t,i, not a shifted reset observation. k=0 measures immediate displacement against the current goal. Other terms measure displacement since recent goals were issued. K_t includes only available offsets within the same episode, and the denominator is the count of valid offsets. This prevents early episode steps from being diluted by nonexistent history. All state and goal arguments are detached: intrinsic reward is data, not a path for backpropagating the worker loss into the manager. Cosine is U(a)^T U(b), with U from the timing slide, so zero vectors yield zero. It scores direction rather than distance. The legacy reward helper ending at pre-action s_t is retained for diagnostics but is not the training path. The terminal action can still earn intrinsic reward using its pre-reset successor.",
"Sources: manager.py::worker_intrinsic_reward_aligned • trainer.py::_apply_intrinsic_reward")
s.eq(r"r^I_{t,i}=\frac{1}{|K_t|}\sum_{k\in K_t}\operatorname{cos}_{\varepsilon}\!\left(s^+_{t,i}-s_{t-k,i},\ g_{t-k,i}\right)",48,146,862,92,30)
s.card(44,279,272,165,"WHAT MOVED?", "s⁺ₜ,ᵢ: the actual successor representation.",TEAL,WHITE,20)
s.card(345,279,272,165,"WHAT WAS ASKED?", "gₜ₋ₖ,ᵢ: a recent goal; k = 0 includes this step.",TEAL,WHITE,20)
s.card(646,279,272,165,"HOW IS IT SCORED?", "Cosine alignment, averaged over valid history Kₜ.",TEAL,WHITE,20)

# 05 / GAE
s=Slide("Advantages / generalized advantage estimation", "GAE assigns credit over time", "Each stream gets its own advantage and return target; standardize only the advantages.",
"q ranges over E (worker extrinsic), I (worker intrinsic), and M (manager extrinsic). Agent and environment indices are suppressed in these equations. v_t^q is the saved rollout critic prediction. r-tilde is the stream reward with a time-limit bootstrap correction: r-tilde_t^q=r_t^q+gamma_q*b_t*V^q(X_t^+), where b_t=1 for a time-limit truncation and zero otherwise. X_t^+ is the actual successor before reset. d_t=1 for termination or truncation, preventing GAE from crossing an episode reset. True terminals have no continuation bootstrap. The saved next value v_(t+1)^q is used on ordinary steps and is the final bootstrap value at rollout end. gamma_q discounts future rewards, lambda controls the GAE trace, delta is the TD residual, A is raw advantage, and R=A+v is the fixed value target. Start the reverse recursion with zero beyond the rollout. Normalize over rollout time independently for each environment and output, using sample standard deviation (ddof=1) and epsilon_A=1e-8. Defaults: gamma_E=gamma_I=0.99, gamma_M=0.99, lambda=0.95. Manager GAE runs at every step; c enters transition scoring, not a separate c-step GAE.",
"Source: mappo.py::compute_gae / ppo_update / manager_update")
s.eq(r"\delta_t^q=\widetilde r_t^q+\gamma_q(1-d_t)v_{t+1}^q-v_t^q",48,130,860,54,29)
s.eq(r"A_t^q=\delta_t^q+\gamma_q\lambda(1-d_t)A_{t+1}^q,\qquad R_t^q=A_t^q+v_t^q",48,201,860,58,28)
s.eq(r"\widehat A_t^q=\frac{A_t^q-\operatorname{mean}_t(A^q)}{\operatorname{std}_t(A^q)+\varepsilon_A}",48,284,545,78,27)
s.text("q = E, I, M\nδ: TD error; A: advantage\nR: value target; v: saved value",610,282,302,19,MUTED)
s.text("r̃: reward + truncation bootstrap   •   d: episode boundary   •   εA = 10⁻⁸",48,398,862,16,MUTED)
s.text("γ = γM = 0.99: discount   •   λ = 0.95: trace decay   •   mean/std over time",48,426,862,16,MUTED)

# 06 / Worker and critic losses
s=Slide("Loss functions / worker and critics", "PPO updates the worker; MSE fits the critics", "PPO reuses stored actions and goals. Each critic fits its own fixed, unnormalized target.",
"rho is the current action probability divided by the rollout action probability. Both use the same observation, stored pooled goal and legal-action mask. phi is the current worker parameter vector and phi_old is the rollout worker. A^W is the extrinsic or mixed advantage from slide 3. The minimum implements the PPO clipped surrogate; epsilon_P=0.2 by default. H is action-distribution entropy, encouraged by subtracting beta*H from the minimized loss; beta=0.01 by default. Angle brackets W denote the mean over active decisions. Fixed actions, goals and advantages do not backpropagate into the manager. Each q in E,I,M has a separate critic V^q with parameters psi_q. R^q is its unnormalized GAE return target. c_q is the value-loss weight, default 0.5, and angle brackets q average that critic's samples (masking inactive per-agent heads). Value regression is ordinary MSE, without value clipping or an extra one-half factor. All critics read X. Worker critics produce one output per agent. The manager critic produces a team output for shared rewards, or per-agent outputs for per-agent learning rewards. The actor and all critics use separate optimizers. These defaults may be overridden by a run's configuration.",
"Sources: mappo.py::ppo_update / manager_update")
s.eq(r"\rho_{t,i}=\frac{\pi_\phi(a_{t,i}\mid o_{t,i},u_{t,i})}{\pi_{\phi_{\rm old}}(a_{t,i}\mid o_{t,i},u_{t,i})}",48,124,862,71,26)
s.eq(r"\mathcal{L}_W=-\left\langle\min\!\left(\rho A^W,\ \operatorname{clip}(\rho,1-\epsilon_P,1+\epsilon_P)A^W\right)\right\rangle_W",48,204,866,46,24)
s.eq(r"{}-\beta\langle\mathcal{H}(\pi_\phi)\rangle_W",82,253,830,35,24)
s.line(48,307,912,307,LINE,1)
s.eq(r"\mathcal{L}_{V^q}=c_q\left\langle\left(V^q_{\psi_q}(X_t)-R_t^q\right)^2\right\rangle_q",48,318,862,56,27)
s.text("ρ: policy ratio   •   εP = 0.2: clipping range   •   β = 0.01: entropy weight",48,389,864,16,MUTED)
s.text("H: entropy   •   Vq: critic   •   ψq: critic parameters   •   cq = 0.5: value-loss weight",48,415,864,16,MUTED)
s.text("⟨·⟩: mean over applicable active samples   •   φ / φold: current / rollout worker",48,441,864,15,MUTED)

# 07 / Manager
s=Slide("Loss functions / manager", "Reinforce goals aligned with useful transitions", "Positive task advantage reinforces alignment; negative advantage discourages it.",
"D is the stabilized cosine between observed c-step displacement and the manager's goal. sg means stop_gradient. The loss is negative normalized extrinsic manager advantage times D, averaged over active samples with a full horizon inside one episode. A scalar team advantage is broadcast to all agents; with per-agent rewards each agent uses its own manager advantage. The final c stored states lack complete future targets and do not contribute, and episode-crossing horizons are masked. The manager has no goal-probability ratio, PPO clipping, or entropy term. It recomputes g(theta) differentiably from the ordered trajectory. The explicit gradient is grad_theta L_M = -mean_M[hat A^M * grad_theta D]. The displacement and advantage are fixed targets, while gradients flow through goal generation. The encoder remains trainable through the goal branch because the manager core consumes s; detach does not freeze the whole encoder. This prevents direct target-rotation gradients but does not guarantee that latent collapse cannot emerge across updates. Its separate critic uses the MSE formula on slide 6, with q=M.",
"Sources: manager.py::transition_cosine • mappo.py::manager_update")
s.eq(r"D_{t,i}=\operatorname{cos}_{\varepsilon}\!\left(\operatorname{sg}[s_{t+c,i}-s_{t,i}],\ g_{t,i}(\theta)\right)",48,128,862,65,29)
s.eq(r"\mathcal{L}_M=-\left\langle\operatorname{sg}[\widehat A^M_{t,i}]\,D_{t,i}\right\rangle_M",48,205,862,65,32,TEAL)
s.text("sg: stop gradient   •   θ: manager parameters   •   ⟨·⟩M: mean over valid, active horizons",48,285,862,16,MUTED)
s.card(44,330,421,116,"FIXED TARGET", "Displacement and advantage are treated as data.",MUTED,WHITE,20)
s.card(497,330,421,116,"TRAINABLE GOAL", "Gradients flow through goal generation and its encoder.",TEAL,WHITE,20)

# 08 / Update order
s=Slide("Updates / one training iteration", "Collect once; update worker, then manager", "Worker: clipped PPO. Manager: task-weighted transition alignment.",
"Collection freezes the networks and saves states, observations, actions, old log probabilities, goals, rewards, values and boundaries. Then build intrinsic rewards if enabled, GAE and return targets. Worker PPO uses shuffled timestep-centric minibatches for six epochs by default, updating the actor and worker critics while holding the manager fixed and reusing stored goals. Do not infer exactly 48 optimizer steps from the nominal minibatch setting: actual partitioning follows the code's sample dimensions. Manager GAE uses saved rollout values. Its critic fits fixed targets for eight full-batch passes, then the manager actor takes one ordered full-batch transition-gradient step. The ordered trajectory is needed for s_t and s_(t+c). Extra manager policy passes have no importance correction, so one is the default. Separate Adam optimizers clip global gradient norm at 0.5, with default learning rates 3e-4. The shared-encoder variant adds a worker encoder gradient to the manager step, but that extension is outside this short base-algorithm presentation. The core message: the manager learns which directional goals accompany useful outcomes; PPO trains the worker to act on task advantage and optional goal-following advantage.",
"Sources: trainer.py::update_fn • mappo.py::ppo_update / manager_update")
steps=[("1","COLLECT","Freeze networks;\nsave rollout data.",MUTED),("2","WORKER PPO","6 epochs; actor\nand worker critics.",TEAL),("3","MANAGER VALUE","8 full-batch passes\non fixed targets.",TEAL),("4","MANAGER GOALS","1 full-batch\ntransition update.",TEAL)]
for j,(num,label,body,col) in enumerate(steps):
    x=44+j*225
    s.text(num,x,144,45,32,col,True)
    s.card(x,209,198,174,label,body,col,WHITE,20)
    if j<3:s.line(x+202,295,x+219,295,col,2,True)
s.text("Adam learning rates: 3 × 10⁻⁴   •   Global gradient-norm clipping: 0.5",48,419,867,18,MUTED)


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def render():
    deck=Presentation()
    deck.slide_width,deck.slide_height=Pt(W),Pt(H)
    deck.core_properties.title="Feudal MAPPO — The Main Algorithm"
    deck.core_properties.subject="Objectives, advantages, losses and update sequence"
    deck.core_properties.author="morphology-learning"
    deck.core_properties.keywords="Feudal, MAPPO, PPO, GAE, intrinsic reward"
    pdf=canvas.Canvas(str(ROOT/'feudal_algorithm.pdf'),pagesize=(W,H))
    pdf.setTitle("Feudal MAPPO — The Main Algorithm")
    pdf.setAuthor("morphology-learning")
    notes=["# Feudal MAPPO — speaker notes", "", "Eight slides · approximately 8–10 minutes · source snapshot: 24 September 2026.", "", "The deck focuses on the main algorithm. Extended notation, boundary handling and caveats are kept in these notes.", ""]
    for s in SLIDES:
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=rgb(WHITE)
        equation_notes="\n\n".join('$$\n'+e+'\n$$' for e in s.equations)
        note=s.notes+"\n\n"+s.source+"\n\nEquations (LaTeX):\n"+equation_notes
        slide.notes_slide.notes_text_frame.text=note
        notes.extend([f"## {s.number}. {s.title}", "", s.notes, "", s.source, "", equation_notes, ""])
        for el in s.elements:
            k=el['kind']
            if k=='rect':
                x,y,w,h=el['x'],el['y'],el['w'],el['h']
                shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if el['radius'] else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
                shape.fill.solid(); shape.fill.fore_color.rgb=rgb(el['fill'])
                if el['stroke']:
                    shape.line.color.rgb=rgb(el['stroke']);shape.line.width=Pt(0.8)
                else:shape.line.fill.background()
                pdf.setFillColor(HexColor('#'+el['fill']))
                if el['stroke']:pdf.setStrokeColor(HexColor('#'+el['stroke']))
                pdf.setLineWidth(0.8)
                if el['radius']:pdf.roundRect(x,H-y-h,w,h,el['radius'],fill=1,stroke=bool(el['stroke']))
                else:pdf.rect(x,H-y-h,w,h,fill=1,stroke=bool(el['stroke']))
            elif k=='text':
                tb=slide.shapes.add_textbox(Pt(el['x']),Pt(el['y']),Pt(el['w']),Pt(el['h']))
                tf=tb.text_frame
                tf.word_wrap=False
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                tf.vertical_anchor=MSO_ANCHOR.TOP
                for j,line in enumerate(el['text'].split('\n')):
                    p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
                    p.text=line
                    p.font.name=FONT;p.font.size=Pt(el['size']);p.font.bold=el['bold'];p.font.color.rgb=rgb(el['color'])
                    p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(el['size']*1.18)
                    p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[el['align']]
                font='Deck-Bold' if el['bold'] else 'Deck'
                pdf.setFont(font,el['size']);pdf.setFillColor(HexColor('#'+el['color']))
                for j,line in enumerate(el['text'].split('\n')):
                    yy=H-el['y']-el['size']*.92-j*el['size']*1.18
                    if el['align']=='center':pdf.drawCentredString(el['x']+el['w']/2,yy,line)
                    elif el['align']=='right':pdf.drawRightString(el['x']+el['w'],yy,line)
                    else:pdf.drawString(el['x'],yy,line)
            elif k=='image':
                slide.shapes.add_picture(str(el['path']),Pt(el['x']),Pt(el['y']),Pt(el['w']),Pt(el['h']))
                pdf.drawImage(str(el['path']),el['x'],H-el['y']-el['h'],el['w'],el['h'],mask='auto')
            elif k=='line':
                line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(el['x1']),Pt(el['y1']),Pt(el['x2']),Pt(el['y2']))
                line.line.color.rgb=rgb(el['color']);line.line.width=Pt(el['width'])
                pdf.setStrokeColor(HexColor('#'+el['color']));pdf.setLineWidth(el['width'])
                pdf.line(el['x1'],H-el['y1'],el['x2'],H-el['y2'])
            elif k=='circle':
                shape=slide.shapes.add_shape(MSO_SHAPE.OVAL,Pt(el['x']),Pt(el['y']),Pt(el['w']),Pt(el['h']))
                shape.fill.solid();shape.fill.fore_color.rgb=rgb(el['fill']);shape.line.fill.background()
                pdf.setFillColor(HexColor('#'+el['fill']))
                pdf.circle(el['x']+el['w']/2,H-el['y']-el['h']/2,el['w']/2,stroke=0,fill=1)
            elif k=='polygon':
                pts=el['points']
                b=slide.shapes.build_freeform(pts[0][0],pts[0][1],scale=12700)
                b.add_line_segments(pts[1:],close=True)
                shape=b.convert_to_shape();shape.fill.solid();shape.fill.fore_color.rgb=rgb(el['fill']);shape.line.fill.background()
                path=pdf.beginPath();path.moveTo(pts[0][0],H-pts[0][1])
                for x,y in pts[1:]:path.lineTo(x,H-y)
                path.close();pdf.setFillColor(HexColor('#'+el['fill']));pdf.drawPath(path,stroke=0,fill=1)
        pdf.showPage()
    pdf.save()
    deck.save(str(ROOT/'feudal_algorithm.pptx'))
    (ROOT/'speaker_notes.md').write_text('\n'.join(notes))
    (ROOT/'README.md').write_text('''# Feudal MAPPO presentation

Eight slides, approximately 8–10 minutes. White background, restrained blue accent,
editable PowerPoint text and diagrams, rendered equations, and speaker notes.

- `feudal_algorithm.pptx`: presentation with embedded speaker notes.
- `feudal_algorithm.pdf`: matching PDF for sharing.
- `speaker_notes.md`: explanations, all equation terms, source references, and LaTeX.
- `build_short_deck.py`: source for the delivered eight-slide deck.
- `build_slides.py`: preserved, earlier long-draft source; not used for the delivered deck.

To rebuild, install `python-pptx`, `reportlab`, `matplotlib`, and `Pillow`, and run
`python build_short_deck.py`. The renderer expects DejaVu Sans in the standard
Linux font directory. Equations are embedded at 240 dpi; edit their LaTeX in the
source to regenerate them. Text and diagrams remain native PowerPoint objects.

Sources are the working-tree implementation under `algorithms/feudal_mappo_jax/`
and the defaults under `conf/algorithm/feudal_mappo_jax.yaml`, inspected on
24 September 2026. Algorithm behavior is explained, not benchmarked.
''')
    print(f'Built {len(SLIDES)} slides: {ROOT / "feudal_algorithm.pptx"}')
    print(f'PDF: {ROOT / "feudal_algorithm.pdf"}')


if __name__=='__main__':
    render()
